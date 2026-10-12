#!/usr/bin/env python3
"""
Convert an LHD classes spreadsheet (xlsx) into ClassList.json
for the RDASH Learning Hub.

Usage:
    python scripts/convert_classlist.py <input_xlsx> [output_json]

If output_json is omitted, defaults to Data/ClassList.json relative
to the repository root.

The spreadsheet must contain the following columns (header names
matched case-insensitively):

    Course, Offering Name, Delivery Mode, Title,
    Start Date, End Date, Start Time, End Time,
    Category, Places Remaining, Offering link

Dates in the spreadsheet may be Excel datetime cells or strings
in YYYY-MM-DD format. They are converted to Excel serial numbers
(the format the web app expects).
"""

import json
import sys
import os
import re
from datetime import datetime, date

try:
    import openpyxl
except ImportError:
    sys.exit("ERROR: openpyxl is required. Install with: pip install openpyxl")

# Excel epoch: December 30, 1899
EXCEL_EPOCH = date(1899, 12, 30)

# Columns we extract from the spreadsheet, in the order they appear
# in the output JSON. The key is the JSON field name; the value is
# the header text to look for (matched case-insensitively).
COLUMN_MAP = {
    "Course": "Course",
    "Offering Name": "Offering Name",
    "Delivery Mode": "Delivery Mode",
    "Title": "Title",
    "Start Date": "Start Date",
    "End Date": "End Date",
    "Start Time": "Start Time",
    "End Time": "End Time",
    "Category": "Category",
    "Places Remaining": "Places Remaining",
    "Offering link": "Offering link",
}

# Required columns — script will error if these are missing
REQUIRED_COLUMNS = {"Course", "Start Date", "End Date"}


def excel_serial(d):
    """Convert a date or datetime to an Excel serial number."""
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, date):
        return (d - EXCEL_EPOCH).days
    return None


def parse_date_cell(value):
    """Parse a date from a spreadsheet cell.

    Handles openpyxl datetime objects, strings (YYYY-MM-DD or DD/MM/YYYY),
    and numeric Excel serials.
    """
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return excel_serial(value)
    if isinstance(value, date):
        return excel_serial(value)

    if isinstance(value, (int, float)):
        # Already an Excel serial
        return int(value)

    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return excel_serial(datetime.strptime(s, fmt).date())
        except ValueError:
            continue

    return None


def find_header_columns(headers):
    """Build a mapping from JSON field name -> column index (0-based)."""
    # Normalise headers for case-insensitive matching
    normalised = {}
    for i, h in enumerate(headers):
        if h is not None:
            normalised[str(h).strip().lower()] = i

    mapping = {}
    for json_field, header_text in COLUMN_MAP.items():
        idx = normalised.get(header_text.lower())
        if idx is not None:
            mapping[json_field] = idx
    return mapping


def convert(xlsx_path, json_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    # Read headers from row 1
    headers = [cell.value for cell in ws[1]]
    col_map = find_header_columns(headers)

    # Check required columns
    missing = REQUIRED_COLUMNS - col_map.keys()
    if missing:
        sys.exit(
            f"ERROR: Spreadsheet is missing required columns: {', '.join(missing)}\n"
            f"Found columns: {', '.join(str(h) for h in headers if h)}"
        )

    records = []
    warnings = []

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        # Skip completely empty rows
        if all(v is None or v == "" for v in row):
            continue

        record = {}
        for json_field, col_idx in col_map.items():
            raw_value = row[col_idx] if col_idx < len(row) else None

            if json_field in ("Start Date", "End Date"):
                record[json_field] = parse_date_cell(raw_value)
            elif json_field == "Places Remaining":
                try:
                    record[json_field] = int(raw_value) if raw_value is not None else 0
                except (ValueError, TypeError):
                    record[json_field] = 0
            else:
                record[json_field] = str(raw_value).strip() if raw_value is not None else ""

        # Validate required fields
        if not record.get("Course"):
            warnings.append(f"Row {row_idx}: missing Course name, skipped")
            continue
        if record.get("Start Date") is None:
            warnings.append(f"Row {row_idx}: could not parse Start Date for '{record['Course']}', skipped")
            continue
        if record.get("End Date") is None:
            record["End Date"] = record["Start Date"]

        # Ensure Offering link exists (default to empty string)
        record.setdefault("Offering link", "")

        records.append(record)

    # Deduplicate by (Course, Start Date)
    seen = set()
    deduped = []
    for r in records:
        key = (r["Course"].strip().lower(), r["Start Date"])
        if key in seen:
            warnings.append(f"Duplicate session skipped: {r['Course']} (Start Date serial {r['Start Date']})")
            continue
        seen.add(key)
        deduped.append(r)

    # Sort by Start Date, then Course name
    deduped.sort(key=lambda r: (r["Start Date"], r["Course"].lower()))

    # Write JSON
    os.makedirs(os.path.dirname(json_path), exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(deduped, f, indent=2, ensure_ascii=False)

    print(f"✓ Converted {len(deduped)} sessions → {json_path}")
    if warnings:
        print(f"\n⚠ {len(warnings)} warning(s):")
        for w in warnings:
            print(f"  - {w}")

    return deduped, warnings


def main():
    if len(sys.argv) < 2:
        sys.exit(
            "Usage: python scripts/convert_classlist.py <input_xlsx> [output_json]\n"
            "  input_xlsx:  Path to the LHD classes spreadsheet\n"
            "  output_json: Path to write ClassList.json (default: Data/ClassList.json)"
        )

    xlsx_path = sys.argv[1]
    if not os.path.isfile(xlsx_path):
        sys.exit(f"ERROR: File not found: {xlsx_path}")

    # Default output path relative to repo root
    repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    json_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        repo_root, "Data", "ClassList.json"
    )

    convert(xlsx_path, json_path)


if __name__ == "__main__":
    main()